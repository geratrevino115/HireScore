import os
import fitz  # PyMuPDF para PDFs
import docx  # python-docx para DOCX

def extract_text_from_pdf(pdf_path):
    """ Extrae texto de un PDF """
    doc = fitz.open(pdf_path)
    return "\n".join([page.get_text() for page in doc])

def extract_text_from_docx(docx_path):
    """ Extrae texto de un DOCX (párrafos y tablas) """
    doc = docx.Document(docx_path)
    text = []

    # Extraer texto de los párrafos
    for para in doc.paragraphs:
        text.append(para.text)

    # Extraer texto de las tablas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text.append(cell.text)

    return "\n".join(text)

def extract_text(file_path):
    """ Detecta el tipo de archivo (PDF/DOCX) y extrae el texto """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"El archivo '{file_path}' no se encontró.")

    ext = os.path.splitext(file_path)[-1].lower()
    
    if ext == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError("Formato no soportado. Solo se permiten archivos PDF y DOCX.")
    
    # Guardar el texto en un archivo de salida
    save_text(file_path, text)
    
    return text

def save_text(file_path, text):
    """ Guarda el texto extraído en un archivo TXT dentro de la carpeta outputs """
    output_dir = "cv_processor/outputs"
    os.makedirs(output_dir, exist_ok=True)  # Crear la carpeta si no existe
    
    # Obtener el nombre base del archivo sin la extensión
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    safe_name = base_name.replace(" ", "_")  # Reemplazar espacios por guiones bajos
    output_path = os.path.join(output_dir, f"{safe_name}.txt")
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)
    
    print(f"Texto guardado en: {output_path}")

if __name__ == "__main__":
    # Para probar directamente este módulo
    file_path = "cv_processor/uploads/Cvejemplo.pdf"
    try:
        text = extract_text(file_path)
        print("\nTexto extraído (primeros 1000 caracteres):\n")
        print(text[:1000])  # Mostrar solo una parte del texto
    except Exception as e:
        print(f"Error: {e}")
